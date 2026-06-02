# 文档加载器

from pathlib import Path

from pypdf import PdfReader
from docx import Document
import markdown
from bs4 import BeautifulSoup

from app.core.logger import logger
from app.rag.doucument_cleaner import DoucumentCleaner


class DocumentLoader:
    def __init__(self, docs_dir: str):
        self.docs_dir = Path(docs_dir)
        self.cleaner = DoucumentCleaner()

    def load_doucuments(self) -> list[dict]:
        """
        返回docs_dir下的所有txt,md,pdf,docx文件
        返回格式：
        [
            {
                "source":"文件名",
                "file_type":"文件类型"
                "content":"文件内容"
            }
        ]
        """

        documents = []

        if not self.docs_dir.exists():
            logger.warning("文档目录不存在：%s", self.docs_dir)
            return documents

        supported_extensions = [".txt", ".md", ".pdf", ".docx"]

        # 循环遍历知识库文件夹下的直接文件/文件夹
        for file_path in self.docs_dir.iterdir():
            # 是文件 且 文件后缀合法
            if file_path.is_file() and file_path.suffix.lower() in supported_extensions:
                try:
                    doc = self.load_single_file(file_path)

                    if doc and doc["content"].strip():
                        documents.append(doc)
                        logger.info("加载文档成功:%s", file_path.name)

                except Exception as e:
                    logger.error("加载文档失败:%s,error=%s", file_path.name, e)

        return documents

    def load_single_file(self, file_path: Path) -> dict:
        suffix = file_path.suffix.lower()
        content: str

        if suffix == ".txt":
            content = self.load_txt(file_path)
        elif suffix == ".md":
            content = self.load_md(file_path)
        elif suffix == ".pdf":
            content = self.load_pdf(file_path)
        elif suffix == ".docx":
            content = self.load_docx(file_path)
        else:
            raise ValueError(f"不支持的文件类型:{suffix}")

        content = self.cleaner.clean(content)
        return {
            "source": file_path.name,
            "file_type": suffix.replace(".", ""),
            "content": content,
        }

    def load_txt(self, file_path: Path) -> str:
        return file_path.read_text(encoding="utf-8")

    def load_md(self, file_path: Path) -> str:
        md_text = file_path.read_text(encoding="utf-8")

        # Markdown转HTML
        html = markdown.markdown(md_text)

        # HTML转纯文本
        soup = BeautifulSoup(html, "html.parser")
        # 删除所有标签 不同HTML元素 用换行隔开
        text = soup.get_text(separator="\n")

        return text

    def load_pdf(self, file_path: Path) -> str:
        # 不是所有pdf都能直接提取文字
        reader = PdfReader(str(file_path))

        texts = []

        # reader.pages[]
        for page_index, page in enumerate(reader.pages):
            page_text = page.extract_text()

            if page_text:
                texts.append(f"\n\n第 {page_index + 1} 页\n{page_text}")

        return "\n".join(texts)

    def load_docx(self, file_path: Path) -> str:
        document = Document(str(file_path))

        texts = []

        # 读取普通段落
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                texts.append(paragraph.text.strip())

        for table in document.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    if cell_text:
                        row_text.append(cell_text)
                if row_text:
                    texts.append(" | ".join(row_text))
        return "\n".join(texts)
